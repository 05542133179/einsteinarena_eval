from __future__ import annotations

import csv
import json
import re
import urllib.parse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from .common import read_json, read_jsonl, write_json_atomic
from .sync import DEFAULT_BASE_URL, request_json


@dataclass(frozen=True)
class Problem:
    problem_id: int
    slug: str
    title: str
    scoring: str


def parse_problems(payload: Any, selected_slugs: list[str]) -> list[Problem]:
    if isinstance(payload, dict):
        payload = payload.get("problems")
    if not isinstance(payload, list):
        raise TypeError("Official /api/problems response must be an array")
    selected = set(selected_slugs)
    by_slug: dict[str, Problem] = {}
    for index, raw in enumerate(payload):
        if not isinstance(raw, dict):
            raise TypeError(f"Problem item {index} is not an object")
        slug = raw.get("slug")
        if not isinstance(slug, str) or slug not in selected:
            continue
        problem_id = raw.get("id")
        title = raw.get("title")
        scoring = raw.get("scoring")
        if not isinstance(problem_id, int) or isinstance(problem_id, bool):
            raise ValueError(f"Invalid problem id for {slug}: {problem_id!r}")
        if not isinstance(title, str) or not title:
            raise ValueError(f"Invalid problem title for {slug}")
        if scoring not in {"maximize", "minimize"}:
            raise ValueError(f"Invalid scoring for {slug}: {scoring!r}")
        by_slug[slug] = Problem(problem_id, slug, title, scoring)
    missing = sorted(selected - set(by_slug))
    if missing:
        raise ValueError(f"Official problem API is missing selected slugs: {missing}")
    return sorted(by_slug.values(), key=lambda problem: problem.title.casefold())


def parse_leaderboard(payload: Any, slug: str) -> list[dict[str, Any]]:
    if not isinstance(payload, list):
        raise TypeError(f"Leaderboard response must be an array: {slug}")
    parsed: list[dict[str, Any]] = []
    for index, raw in enumerate(payload):
        if not isinstance(raw, dict):
            raise TypeError(f"Leaderboard item {index} is not an object: {slug}")
        rank = raw.get("rank")
        agent = raw.get("agentName")
        score = raw.get("bestScore")
        submissions = raw.get("submissions")
        if not isinstance(rank, int) or rank != index + 1:
            raise ValueError(f"Non-contiguous official rank for {slug}: {rank!r}")
        if not isinstance(agent, str) or not agent:
            raise ValueError(f"Missing agentName for {slug} rank {rank}")
        if (
            not isinstance(score, (int, float))
            or isinstance(score, bool)
        ):
            raise ValueError(f"Invalid score for {slug} rank {rank}: {score!r}")
        if not isinstance(submissions, int) or isinstance(submissions, bool):
            raise ValueError(f"Invalid submissions for {slug} rank {rank}")
        parsed.append(
            {
                "rank": rank,
                "agent": agent,
                "score": float(score),
                "submissions": submissions,
            }
        )
    return parsed


def _load_local_best(
    path: Path | None,
    problems: list[Problem],
) -> dict[str, float]:
    if path is None:
        return {}
    scoring = {problem.slug: problem.scoring for problem in problems}
    best: dict[str, float] = {}
    for row in read_jsonl(path):
        slug = str(row.get("slug") or "")
        score = row.get("raw_score")
        if (
            slug not in scoring
            or row.get("valid") is not True
            or not isinstance(score, (int, float))
            or isinstance(score, bool)
        ):
            continue
        numeric = float(score)
        if slug not in best:
            best[slug] = numeric
        elif scoring[slug] == "maximize":
            best[slug] = max(best[slug], numeric)
        else:
            best[slug] = min(best[slug], numeric)
    return best


def _virtual_rank(
    problem: Problem,
    score: float,
    official: list[dict[str, Any]],
) -> int:
    if problem.scoring == "maximize":
        ahead = sum(float(row["score"]) >= score for row in official)
    else:
        ahead = sum(float(row["score"]) <= score for row in official)
    return ahead + 1


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0]) if rows else ["problem_slug"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def _write_xlsx(
    path: Path,
    overview: list[dict[str, Any]],
    long_rows: list[dict[str, Any]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    for index, (name, rows) in enumerate(
        (("Overview", overview), ("Leaderboard", long_rows))
    ):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        headers = list(rows[0]) if rows else ["problem_slug"]
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="355C7D")
            cell.alignment = Alignment(horizontal="center")
        for column, header in enumerate(headers, start=1):
            width = max(
                len(str(header)),
                *(
                    len(str(row.get(header) or ""))
                    for row in rows[:200]
                ),
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max(width + 2, 12), 42
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def pull_leaderboard(
    data_dir: Path,
    output_root: Path,
    *,
    base_url: str = DEFAULT_BASE_URL,
    top_k: int = 10,
    leaderboard_limit: int = 100,
    timeout: float = 30,
    retries: int = 3,
    local_results: Path | None = None,
    local_label: str = "local",
    snapshot: str | None = None,
    problems_payload: Any | None = None,
    leaderboard_payloads: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if not 1 <= top_k <= leaderboard_limit <= 100:
        raise ValueError("Require 1 <= top_k <= leaderboard_limit <= 100")
    manifest = read_json(data_dir / "manifest.json")
    selected = manifest.get("open_problem_slugs")
    if not isinstance(selected, list) or not selected:
        raise ValueError("Dataset manifest has no open_problem_slugs")
    problems_url = f"{base_url.rstrip('/')}/api/problems"
    raw_problems = (
        problems_payload
        if problems_payload is not None
        else request_json(problems_url, timeout=timeout, retries=retries)
    )
    problems = parse_problems(raw_problems, [str(value) for value in selected])
    parsed: dict[str, list[dict[str, Any]]] = {}
    raw_leaderboards: dict[str, Any] = {}
    urls: dict[str, str] = {}
    for problem in problems:
        url = (
            f"{base_url.rstrip('/')}/api/leaderboard?"
            + urllib.parse.urlencode(
                {"problem_id": problem.problem_id, "limit": leaderboard_limit}
            )
        )
        payload = (
            leaderboard_payloads[problem.slug]
            if leaderboard_payloads is not None
            else request_json(url, timeout=timeout, retries=retries)
        )
        urls[problem.slug] = url
        raw_leaderboards[problem.slug] = payload
        parsed[problem.slug] = parse_leaderboard(payload, problem.slug)
        print(
            f"Fetched leaderboard: {problem.slug} "
            f"({len(parsed[problem.slug])} entries)",
            flush=True,
        )
    local_best = _load_local_best(local_results, problems)
    overview: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    for problem in problems:
        official = parsed[problem.slug]
        leader = official[0] if official else None
        local_score = local_best.get(problem.slug)
        overview.append(
            {
                "problem_slug": problem.slug,
                "problem_title": problem.title,
                "scoring": problem.scoring,
                "official_best_agent": leader["agent"] if leader else "",
                "official_best_score": leader["score"] if leader else "",
                "local_label": local_label if local_score is not None else "",
                "local_best_score": (
                    local_score if local_score is not None else ""
                ),
                "local_virtual_rank": (
                    _virtual_rank(problem, local_score, official)
                    if local_score is not None
                    else ""
                ),
            }
        )
        for row in official[:top_k]:
            long_rows.append(
                {
                    "problem_slug": problem.slug,
                    "problem_title": problem.title,
                    "scoring": problem.scoring,
                    "rank": row["rank"],
                    "agent": row["agent"],
                    "score": row["score"],
                    "submissions": row["submissions"],
                    "source": "official",
                }
            )
        if local_score is not None:
            long_rows.append(
                {
                    "problem_slug": problem.slug,
                    "problem_title": problem.title,
                    "scoring": problem.scoring,
                    "rank": _virtual_rank(problem, local_score, official),
                    "agent": local_label,
                    "score": local_score,
                    "submissions": "",
                    "source": "local_eval",
                }
            )
    snapshot_name = snapshot or datetime.now(
        ZoneInfo("Asia/Hong_Kong")
    ).strftime("%Y-%m-%d_%H%M%S")
    if not re.fullmatch(r"[0-9A-Za-z_.-]+", snapshot_name):
        raise ValueError("Snapshot may contain only letters, digits, . _ -")
    output_dir = output_root.resolve() / snapshot_name
    if output_dir.exists() and any(output_dir.iterdir()):
        raise FileExistsError(f"Output directory is not empty: {output_dir}")
    overview_path = output_dir / "overview.csv"
    long_path = output_dir / "leaderboard.csv"
    xlsx_path = output_dir / "leaderboard.xlsx"
    raw_path = output_dir / "raw.json"
    _write_csv(overview_path, overview)
    _write_csv(long_path, long_rows)
    _write_xlsx(xlsx_path, overview, long_rows)
    write_json_atomic(
        raw_path,
        {
            "snapshot": snapshot_name,
            "problems_url": problems_url,
            "leaderboard_urls": urls,
            "problems": raw_problems,
            "leaderboards": raw_leaderboards,
            "local_results": str(local_results.resolve()) if local_results else None,
        },
    )
    return {
        "snapshot": snapshot_name,
        "problems": len(problems),
        "local_problems": len(local_best),
        "output_dir": str(output_dir),
        "overview_csv": str(overview_path),
        "leaderboard_csv": str(long_path),
        "xlsx": str(xlsx_path),
        "raw_json": str(raw_path),
    }
